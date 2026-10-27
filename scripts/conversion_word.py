"""PDF coordinates -> editable Word drawings, without paragraph reflow or overlays."""
import copy
import io
import math
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Pt
from fontTools.ttLib import TTFont
from lxml import etree as ET
from pptx import Presentation

WPS = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'


def node(tag, attrs=None):
    element = OxmlElement(tag)
    for key, value in (attrs or {}).items():
        element.set(qn(key) if ':' in key else key, str(value))
    return element


def anchor(x, y, width, height, identifier, order=0):
    emu = lambda value: round(value * 12700)
    element = node('wp:anchor', dict(distT=0, distB=0, distL=0, distR=0, simplePos=0,
                                    relativeHeight=max(0, order), behindDoc=0, locked=0, layoutInCell=0, allowOverlap=1))
    element.append(node('wp:simplePos', dict(x=0, y=0)))
    for tag, relative, position in (('positionH', 'page', x), ('positionV', 'page', y)):
        axis = node(f'wp:{tag}', dict(relativeFrom=relative))
        offset = node('wp:posOffset')
        offset.text = str(emu(position))
        axis.append(offset)
        element.append(axis)
    element.append(node('wp:extent', dict(cx=max(1, emu(width)), cy=max(1, emu(height)))))
    element.append(node('wp:wrapNone'))
    element.append(node('wp:docPr', dict(id=identifier, name=f'PDF object {identifier}')))
    element.append(node('wp:cNvGraphicFramePr'))
    return element


def attach(paragraph, element):
    run = paragraph.add_run()
    drawing = node('w:drawing')
    drawing.append(element)
    run._r.append(drawing)


def shape_anchor(x, y, width, height, identifier, order=0, properties=None):
    element = anchor(x, y, width, height, identifier, order)
    graphic = node('a:graphic')
    data = node('a:graphicData', dict(uri=WPS))
    shape = ET.Element(f'{{{WPS}}}wsp', nsmap={'wps': WPS})
    shape.append(ET.Element(f'{{{WPS}}}cNvSpPr'))
    sp = ET.SubElement(shape, f'{{{WPS}}}spPr')
    if properties is None:
        transform = node('a:xfrm')
        transform.append(node('a:off', dict(x=0, y=0)))
        transform.append(node('a:ext', dict(cx=round(width*12700), cy=round(height*12700))))
        sp.append(transform)
        geometry = node('a:prstGeom', dict(prst='rect'))
        geometry.append(node('a:avLst'))
        sp.append(geometry)
        sp.append(node('a:noFill'))
        stroke = node('a:ln')
        stroke.append(node('a:noFill'))
        sp.append(stroke)
    else:
        for child in properties:
            sp.append(copy.deepcopy(child))
        offset = sp.find(f'{{{A}}}xfrm/{{{A}}}off')
        offset.set('x', '0')
        offset.set('y', '0')
    data.append(shape)
    graphic.append(data)
    element.append(graphic)
    return element, shape


def font_metrics():
    result = {}
    for path in (Path(__file__).resolve().parents[1] / 'backend/fonts').rglob('*.ttf'):
        font = TTFont(path)
        names = [record.toUnicode() for record in font['name'].names if record.nameID in (1, 16)]
        flags = font['OS/2'].fsSelection if 'OS/2' in font else 0
        for name in names:
            result[(name, bool(flags & 32), bool(flags & 1))] = font
    return result


def build_layout_docx(source, output, text_lines, font_name, rgb, slide_paths):
    document = Document()
    for setting in document.settings._element.xpath('//w:compatSetting[@w:name="compatibilityMode"]'):
        setting.set(qn('w:val'), '15')
    document.styles['Normal'].paragraph_format.space_before = Pt(0)
    document.styles['Normal'].paragraph_format.space_after = Pt(0)
    document.styles['Normal'].font.size = Pt(1)
    fonts = font_metrics()
    identifier = 1
    with fitz.open(source) as pdf:
        for page in pdf:
            section = document.sections[0] if page.number == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
            section.page_width, section.page_height = Pt(page.rect.width), Pt(page.rect.height)
            section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(0)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = Pt(1)
            # Keep artwork as separate editable vector shapes, not a page image.
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide_paths(page, slide, lambda point: (round(point.x*12700), round(point.y*12700)))
            for shape in slide.shapes:
                element, wsp = shape_anchor(shape.left/12700, shape.top/12700, max(shape.width/12700, 0.01),
                                            max(shape.height/12700, 0.01), identifier, properties=shape._element.spPr)
                element.set('behindDoc', '1')
                ET.SubElement(wsp, f'{{{WPS}}}bodyPr')
                attach(paragraph, element)
                identifier += 1
            for block in page.get_text('dict')['blocks']:
                if block['type'] != 1:
                    continue
                box = fitz.Rect(block['bbox'])
                picture = paragraph.add_run().add_picture(io.BytesIO(block['image']), width=Pt(box.width), height=Pt(box.height))
                inline = picture._inline
                element = anchor(box.x0, box.y0, box.width, box.height, identifier, 100000)
                element.append(copy.deepcopy(inline.find(qn('a:graphic'))))
                inline.getparent().replace(inline, element)
                identifier += 1
            links = [link for link in page.get_links() if link.get('uri', '').startswith(('https://', 'http://', 'mailto:'))]
            for line in text_lines(page):
                direction = line.get('dir', (1, 0))
                vertical = abs(direction[1]) > 0.5
                for span in line['spans']:
                    text = span['text']
                    if not text.strip():
                        continue
                    box = fitz.Rect(span['bbox'])
                    name = font_name(span['font'])
                    bold, italic = bool(span['flags'] & 16), bool(span['flags'] & 2)
                    font = fonts.get((name, bold, italic)) or fonts.get((name, False, False))
                    size = span['size']
                    # Anchor each source run once. PDF x coordinates preserve
                    # justification, mixed fonts, table columns and maths runs.
                    x = box.x0
                    y = span['origin'][1] - size
                    width, height = max(box.width + size*2, 1), max(size*2, box.height)
                    if vertical:
                        x, y, width, height = box.x0, box.y0, max(box.width, size*1.5), box.height+size
                    # Word clips italic overhangs and some maths faces more
                    # aggressively than PDF viewers. Pad the outer shape and
                    # compensate its origin so the text stays in place.
                    padding = max(2, size * 0.4)
                    x, y = x - padding, y - padding
                    width, height = width + padding*2, height + padding*2
                    element, shape = shape_anchor(x, y, width, height, identifier, 200000)
                    shape[0].set('txBox', '1')
                    textbox = ET.SubElement(shape, f'{{{WPS}}}txbx')
                    content = ET.SubElement(textbox, qn('w:txbxContent'))
                    p = node('w:p')
                    ppr = node('w:pPr')
                    ppr.append(node('w:spacing', {'w:before': 0, 'w:after': 0, 'w:line': round(size*24), 'w:lineRule': 'atLeast'}))
                    p.append(ppr)
                    run = node('w:r')
                    props = node('w:rPr')
                    props.append(node('w:rFonts', {'w:ascii': name, 'w:hAnsi': name, 'w:eastAsia': name, 'w:cs': name}))
                    props.append(node('w:b', {'w:val': int(bold)}))
                    props.append(node('w:i', {'w:val': int(italic)}))
                    props.append(node('w:sz', {'w:val': round(size*2)}))
                    props.append(node('w:color', {'w:val': rgb(span['color'])}))
                    # Scale usable font metrics to the PDF run width. This also
                    # prevents a fallback face from pushing adjacent runs apart.
                    if font and not vertical:
                        cmap = font.getBestCmap() or {}
                        advance = sum(font['hmtx'][cmap[ord(char)]][0] for char in text if ord(char) in cmap)
                        if advance:
                            ratio = box.width / (advance/font['head'].unitsPerEm*size)
                            props.append(node('w:w', {'w:val': max(1, min(600, round(ratio*100)))}))
                    run.append(props)
                    value = node('w:t', {'xml:space': 'preserve'})
                    value.text = text
                    run.append(value)
                    link = next((link for link in links if fitz.Rect(link['from']).intersects(box)), None)
                    if link:
                        relationship = document.part.relate_to(link['uri'], RT.HYPERLINK, is_external=True)
                        hyperlink = node('w:hyperlink', {'r:id': relationship})
                        hyperlink.append(run)
                        p.append(hyperlink)
                    else:
                        p.append(run)
                    content.append(p)
                    body = ET.SubElement(shape, f'{{{WPS}}}bodyPr', lIns=str(round(padding*12700)), tIns=str(round(padding*12700)),
                                         rIns=str(round(padding*12700)), bIns=str(round(padding*12700)), wrap='none',
                                         horzOverflow='overflow', vertOverflow='overflow',
                                         vert='vert' if direction[1] > 0.5 else 'vert270' if vertical else 'horz')
                    # Grow the shape for unexpected font metrics; never shrink
                    # or crop the text to make it fit its original PDF box.
                    body.append(node('a:spAutoFit'))
                    attach(paragraph, element)
                    identifier += 1
        document.save(output)
        from conversion_fonts import embed_word_fonts
        embed_word_fonts(pdf, output)
