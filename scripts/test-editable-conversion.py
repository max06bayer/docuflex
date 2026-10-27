#!/usr/bin/env python3
"""Regression tests for content editability, layout and hostile/edge input.

Run with the staged desktop Python. No network or Office installation required.
Set TESSERACT_BIN and TESSDATA_PREFIX to test scans with bundled OCR.
"""
import importlib.util
import io
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile
from lxml import etree as ET

import pymupdf as fitz
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('editable', ROOT / 'scripts/build-converted-office.py')
engine = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(engine)


def fixture(path):
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((48, 60), 'Docuflex Editable Conversion', fontsize=22, fontname='hebo', color=(0.1, 0.2, 0.5))
    page.insert_text((48, 95), 'Real text, not a page screenshot.', fontsize=12)
    page.insert_text((48, 118), 'Bold and italic retain their styles.', fontsize=12, fontname='tibi')
    font = ROOT / 'backend/fonts/inter-variable-normal.ttf'
    page.insert_font(fontname='inter', fontfile=str(font))
    page.insert_text((48, 145), 'Grüße, München - café €', fontsize=12, fontname='inter')
    page.draw_rect((48, 180, 500, 276), color=(0.2, 0.2, 0.2), fill=None)
    for x in (250, 375):
        page.draw_line((x, 180), (x, 276))
    for y in (212, 244):
        page.draw_line((48, y), (500, y))
    for y, values in zip((201, 233, 265), [('Item', 'Quantity', 'Price'), ('Coffee', '2', '12.50'), ('Tea', '3', '7.25')]):
        for x, value in zip((58, 260, 385), values):
            page.insert_text((x, y), value, fontsize=12)
    for y, alias, filename, text in ((440, 'libregular', 'LiberationSans-Regular.ttf', 'Embedded regular font'),
                                     (465, 'libbold', 'LiberationSans-Bold.ttf', 'Embedded bold font')):
        page.insert_font(fontname=alias, fontfile=str(ROOT / 'backend/fonts/fallback' / filename))
        page.insert_text((48, y), text, fontsize=12, fontname=alias)
    page.insert_text((48, 310), '=1+1', fontsize=12)
    page.insert_text((250, 310), '00123', fontsize=12)
    picture = Image.new('RGB', (80, 60), '#3060c0')
    output = io.BytesIO()
    picture.save(output, format='PNG')
    page.insert_image((48, 350, 128, 410), stream=output.getvalue())
    page.draw_bezier((200, 390), (240, 320), (300, 450), (350, 360), color=(0.8, 0.1, 0.1))
    page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(48, 78, 260, 99), 'uri': 'https://docuflex.eu'})
    page = doc.new_page(width=792, height=612)
    page.insert_text((48, 60), 'Landscape second page', fontsize=18)
    page.set_rotation(90)
    doc.save(path)
    doc.close()


class EditableConversion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix='docuflex-conversion-test-')
        cls.directory = Path(cls.temp.name)
        cls.pdf = cls.directory / 'fixture.pdf'
        fixture(cls.pdf)
        for fmt in ('docx', 'pptx', 'xlsx'):
            subprocess.run([sys.executable, str(ROOT / 'scripts/build-converted-office.py'), fmt,
                            str(cls.directory / f'output.{fmt}'), str(cls.pdf)], check=True, capture_output=True)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_word_has_editable_text_vectors_and_styles(self):
        doc = Document(self.directory / 'output.docx')
        text = ' '.join(doc._element.xpath('//w:t/text()'))
        self.assertIn('Real text', text)
        self.assertIn('München', text)
        self.assertIn('Landscape second page', text)
        self.assertIn('Coffee', text)
        self.assertEqual(text.count('Real text'), 1)
        self.assertTrue(doc._element.xpath('//w:txbxContent'))
        self.assertTrue(doc._element.xpath('//w:hyperlink'))
        self.assertTrue(doc._element.xpath('//*[local-name()="spAutoFit"]'))
        self.assertFalse(doc._element.xpath('//w:txbxContent//w:spacing[@w:lineRule="exact"]'))
        self.assertTrue(doc._element.xpath('//*[local-name()="bodyPr" and @horzOverflow="overflow"]'))
        self.assertTrue(doc._element.xpath('//w:rPr/w:b[@w:val="1"]'))
        self.assertTrue(doc._element.xpath('//w:rPr/w:i[@w:val="1"]'))
        self.assertEqual(len(doc._element.xpath('//*[local-name()="pic"]')), 1)

    def test_slides_have_editable_text_vectors_and_only_source_image(self):
        ppt = Presentation(self.directory / 'output.pptx')
        self.assertEqual(len(ppt.slides), 2)
        text = '\n'.join(shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('München', text)
        self.assertIn('Landscape second page', text)
        self.assertEqual(sum(shape.shape_type == 13 for slide in ppt.slides for shape in slide.shapes), 1)
        self.assertTrue(any(shape.shape_type == 5 for shape in ppt.slides[0].shapes))
        self.assertTrue(any(run.hyperlink.address for shape in ppt.slides[0].shapes if shape.has_text_frame
                            for paragraph in shape.text_frame.paragraphs for run in paragraph.runs))

    def test_slides_preserve_bleed_images_alpha_and_text_paint_order(self):
        pdf = self.directory / 'masked-images.pdf'
        output = self.directory / 'masked-images.pptx'
        document = fitz.open()
        page = document.new_page(width=640, height=360)
        page.draw_rect(page.rect, color=None, fill=(0.6,0.2,0.7))
        def png(image):
            data = io.BytesIO()
            image.save(data, format='PNG')
            return data.getvalue()
        # PowerPoint-generated PDFs can put the background a tiny fraction of
        # a point outside MediaBox. TextPage's default clipping drops it.
        page.insert_image((0,-0.00006,640,359.99994), stream=png(Image.new('RGB',(640,360),(40,100,160))), keep_proportion=False)
        logo = Image.new('RGBA',(120,60),(0,0,0,0))
        ImageDraw.Draw(logo).rectangle((20,10,100,50), fill=(255,255,255,255))
        page.insert_image((500,20,620,80), stream=png(logo))
        shadow = Image.new('RGBA',(460,75),(0,0,0,35))
        page.insert_image((35,95,495,170), stream=png(shadow))
        page.insert_text((50,140),'First title line',fontsize=36,color=(1,1,1))
        page.insert_image((35,130,495,205), stream=png(shadow))
        page.insert_text((50,178),'Second title line',fontsize=36,color=(1,1,1))
        document.save(pdf)
        document.close()
        subprocess.run([sys.executable, str(ROOT / 'scripts/build-converted-office.py'), 'pptx',str(output),str(pdf)],check=True,capture_output=True)
        presentation=Presentation(output)
        self.assertEqual(sum(shape.shape_type==13 for shape in presentation.slides[0].shapes),4)
        self.assertIn('Second title line',' '.join(shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame))
        with ZipFile(output) as package:
            images=[Image.open(io.BytesIO(package.read(name))) for name in package.namelist() if name.startswith('ppt/media/')]
            self.assertTrue(any(image.mode=='RGBA' and image.getextrema()[3][0]==0 for image in images))
            fonts=[package.read(name) for name in package.namelist() if name.endswith('.fntdata')]
            self.assertTrue(fonts)
            for data in fonts:
                self.assertEqual(int.from_bytes(data[:4],'little'),len(data))
                self.assertEqual(data[34:36],b'LP')
        if os.environ.get('DOCUMENT_CONVERTER_BIN'):
            directory=self.directory/'render-images'
            directory.mkdir()
            subprocess.run([os.environ['DOCUMENT_CONVERTER_BIN'],f'-env:UserInstallation={(directory/"profile").as_uri()}',
                            '--headless','--convert-to','pdf','--outdir',str(directory),str(output)],check=True,capture_output=True,timeout=60)
            with fitz.open(directory/'masked-images.pdf') as rendered:
                pix=rendered[0].get_pixmap()
                # Photo survived instead of exposing the purple base rectangle.
                self.assertLess(sum(abs(a-b) for a,b in zip(pix.pixel(10,10)[:3],(40,100,160))),15)
                # Transparent corners show the photo rather than solid black.
                self.assertLess(sum(abs(a-b) for a,b in zip(pix.pixel(505,25)[:3],(40,100,160))),15)
                pixels=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
                for box in ((45,110,350,145),(45,148,390,183)):
                    white=sum(all(value>220 for value in color) for color in pixels.crop(box).getdata())
                    self.assertGreater(white,100,'A title line was hidden by a later shadow image.')

    def test_spreadsheet_has_cells_numbers_and_no_formulas(self):
        book = load_workbook(self.directory / 'output.xlsx')
        values = [cell.value for sheet in book for row in sheet for cell in row]
        self.assertIn('Coffee', values)
        self.assertIn(12.5, values)
        price = next(cell for sheet in book for row in sheet for cell in row if cell.value == 12.5)
        self.assertEqual(price.number_format, '0.00')
        self.assertIn('00123', values)
        self.assertTrue(any('=1+1' in str(value) for value in values))
        self.assertFalse(any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row))
        self.assertEqual(sum(len(sheet._images) for sheet in book), 1)

    def test_spreadsheet_preserves_numeric_precision(self):
        from openpyxl import Workbook
        sheet = Workbook().active
        for row, value in enumerate(('0.1234567890123456789', '1234567890123456', '00123', '12.50', '-3', '0'), 1):
            engine.write_cell(sheet.cell(row, 1), value)
        self.assertEqual([sheet.cell(row, 1).value for row in range(1, 7)],
                         ['0.1234567890123456789', '1234567890123456', '00123', 12.5, -3, 0])
        self.assertEqual(sheet['A4'].number_format, '0.00')

    def test_word_embeds_reusable_font(self):
        with ZipFile(self.directory / 'output.docx') as package:
            fonts = [name for name in package.namelist() if name.endswith('.odttf')]
            self.assertGreaterEqual(len(fonts), 1)
            table = ET.fromstring(package.read('word/fontTable.xml'))
            self.assertTrue(table.xpath('//*[local-name()="embedRegular"]'))
            self.assertTrue(table.xpath('//*[local-name()="embedBold"]'))

    @unittest.skipUnless(os.environ.get('DOCUMENT_CONVERTER_BIN'), 'Set Office binary to verify rendered layout')
    def test_office_roundtrip_preserves_all_text_and_rotated_page(self):
        for fmt in ('docx', 'pptx'):
            directory = self.directory / f'render-{fmt}'
            directory.mkdir()
            subprocess.run([os.environ['DOCUMENT_CONVERTER_BIN'],
                            f'-env:UserInstallation={(directory / "profile").as_uri()}',
                            '--headless', '--convert-to', 'pdf', '--outdir', str(directory),
                            str(self.directory / f'output.{fmt}')], check=True, capture_output=True, timeout=60)
            with fitz.open(directory / 'output.pdf') as rendered:
                self.assertEqual(len(rendered), 2)
                text = ' '.join(page.get_text() for page in rendered)
                for expected in ('Real text', 'München', 'Coffee', '12.50', 'Landscape second page'):
                    self.assertIn(expected, text, f'{fmt} lost text when rendered')

    def test_xml_controls_and_tex_delimiters(self):
        self.assertEqual(engine.xml_text('A\x00B\x01C\ufffeD'), 'A\ufffdB\ufffdC\ufffdD')
        self.assertEqual(engine.xml_text('Grüße\t\n€'), 'Grüße\t\n€')
        self.assertEqual(engine.glyph_unicode('parenleftbig'), '(')
        self.assertEqual(engine.glyph_unicode('parenrightBigg'), ')')
        self.assertEqual(engine.glyph_unicode('bracerightBigg'), '}')

    def test_locked_document_is_rejected(self):
        locked = self.directory / 'locked.pdf'
        with fitz.open(self.pdf) as doc:
            doc.save(locked, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='secret', owner_pw='owner')
        result = subprocess.run([sys.executable, str(ROOT / 'scripts/build-converted-office.py'), 'docx',
                                 str(self.directory / 'locked.docx'), str(locked)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn('Unlock the password-protected PDF', result.stderr)

    @unittest.skipUnless(os.environ.get('TESSERACT_BIN'), 'Set bundled OCR environment to test scans')
    def test_scan_is_real_ocr_text(self):
        image = Image.new('RGB', (1400, 500), 'white')
        font = ImageFont.truetype(str(ROOT / 'backend/fonts/inter-variable-normal.ttf'), 70)
        ImageDraw.Draw(image).text((50, 90), 'DOCUFLEX LOCAL SCAN', fill='black', font=font)
        scan = self.directory / 'scan.pdf'
        image.save(scan, resolution=150)
        for fmt in ('docx', 'pptx', 'xlsx'):
            target = self.directory / f'scan.{fmt}'
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/build-converted-office.py'), fmt,
                                     str(target), str(scan)], check=True, capture_output=True, text=True)
            self.assertIn('"ocrPages": 1', result.stdout)
            if fmt == 'docx':
                text = ' '.join(Document(target)._element.xpath('//w:t/text()'))
            elif fmt == 'pptx':
                text = ' '.join(shape.text for slide in Presentation(target).slides for shape in slide.shapes if shape.has_text_frame)
            else:
                text = ' '.join(str(cell.value or '') for sheet in load_workbook(target) for row in sheet for cell in row)
            self.assertIn('DOCUFLEX', text)


if __name__ == '__main__':
    unittest.main(verbosity=2)
